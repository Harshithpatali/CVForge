from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from app.core.config import settings

OUT=Path(settings.storage_dir); OUT.mkdir(parents=True,exist_ok=True)

def render_docx(resume: dict, artifact_id: int):
    path=OUT/f"resume_{artifact_id}.docx"; doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(9.5)
    p=doc.add_paragraph(); p.alignment=1; r=p.add_run(resume["name"]); r.bold=True; r.font.size=Pt(17)
    p=doc.add_paragraph(); p.alignment=1; p.add_run(resume["contact_line"]).font.size=Pt(8.5)
    p=doc.add_paragraph(); p.alignment=1; r=p.add_run(resume["headline"]); r.bold=True; r.font.size=Pt(10.5)
    def heading(t):
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2); r=p.add_run(t.upper()); r.bold=True; r.font.size=Pt(10)
    heading("Summary"); doc.add_paragraph(resume["summary"])
    heading("Skills"); doc.add_paragraph(" • ".join(resume["skills"]))
    if resume["experience"]:
        heading("Experience")
        for e in resume["experience"]:
            p=doc.add_paragraph(); r=p.add_run(f"{e['title']} | {e['company']}"); r.bold=True; p.add_run(f"  {e['dates']}")
            if e.get("location"): p.add_run(f" | {e['location']}")
            for b in e["bullets"]: doc.add_paragraph(b,style="List Bullet")
    if resume["projects"]:
        heading("Projects")
        for pjt in resume["projects"]:
            p=doc.add_paragraph(); r=p.add_run(pjt["name"]); r.bold=True
            if pjt.get("technologies"): p.add_run(f" — {', '.join(pjt['technologies'])}")
            for b in pjt["bullets"]: doc.add_paragraph(b,style="List Bullet")
    if resume["education"]:
        heading("Education")
        for e in resume["education"]:
            p=doc.add_paragraph(); r=p.add_run(f"{e['degree']} {e.get('field','')}"); r.bold=True; p.add_run(f" | {e['institution']} | {e.get('dates','')}")
    if resume["certifications"]:
        heading("Certifications")
        for c in resume["certifications"]: doc.add_paragraph(c,style="List Bullet")
    doc.save(path); return str(path)

def render_pdf(resume: dict, artifact_id: int):
    path=OUT/f"resume_{artifact_id}.pdf"; styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="Name2",parent=styles["Heading1"],alignment=TA_CENTER,fontSize=17,leading=19)); styles.add(ParagraphStyle(name="Center2",parent=styles["Normal"],alignment=TA_CENTER,fontSize=8.5,leading=10)); styles.add(ParagraphStyle(name="H2x",parent=styles["Heading2"],fontSize=10,leading=12,spaceBefore=7,spaceAfter=3)); styles.add(ParagraphStyle(name="Bodyx",parent=styles["BodyText"],fontSize=9,leading=11,spaceAfter=2))
    doc=SimpleDocTemplate(str(path),pagesize=LETTER,rightMargin=42,leftMargin=42,topMargin=35,bottomMargin=35); story=[Paragraph(resume["name"],styles["Name2"]),Paragraph(resume["contact_line"],styles["Center2"]),Paragraph(resume["headline"],styles["Center2"]),Spacer(1,5)]
    def add_section(title): story.extend([Paragraph(title.upper(),styles["H2x"])])
    add_section("Summary"); story.append(Paragraph(resume["summary"],styles["Bodyx"])); add_section("Skills"); story.append(Paragraph(" • ".join(resume["skills"]),styles["Bodyx"]))
    if resume["experience"]:
        add_section("Experience")
        for e in resume["experience"]:
            story.append(Paragraph(f"<b>{e['title']} | {e['company']}</b> — {e['dates']}",styles["Bodyx"])); story += [Paragraph(f"• {b}",styles["Bodyx"]) for b in e["bullets"]]
    if resume["projects"]:
        add_section("Projects")
        for e in resume["projects"]:
            story.append(Paragraph(f"<b>{e['name']}</b> — {', '.join(e.get('technologies',[]))}",styles["Bodyx"])); story += [Paragraph(f"• {b}",styles["Bodyx"]) for b in e["bullets"]]
    if resume["education"]:
        add_section("Education")
        for e in resume["education"]: story.append(Paragraph(f"<b>{e['degree']} {e.get('field','')}</b> | {e['institution']} | {e.get('dates','')}",styles["Bodyx"]))
    if resume["certifications"]:
        add_section("Certifications"); story += [Paragraph(f"• {x}",styles["Bodyx"]) for x in resume["certifications"]]
    doc.build(story); return str(path)
